"""Build the public executive resume as an accessible three-page PDF.

The document follows the approved V27 visual hierarchy and Google XYZ-style
achievement writing. Private contact data is intentionally excluded.
"""

from pathlib import Path
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output/pdf/Joseph_Hanna_Executive_Resume.pdf"
PUBLIC_OUT = ROOT / "public/resume/Joseph_Hanna_Executive_Resume.pdf"

BLUE = "1F4E79"
INK = "20202D"
MUTED = "626262"

CORE_SKILLS = [
    "Enterprise IT Strategy & Roadmaps",
    "Change Management & Workforce Planning",
    "Digital Transformation",
    "Vendor & MSP Governance",
    "Cloud & Hybrid Infrastructure",
    "Executive Partnership & Reporting",
    "Cybersecurity Strategy & Resilience",
    "Business Continuity & Disaster Recovery",
    "ERP & Enterprise Applications",
    "IT Service Management & ITIL",
    "IT Governance, Risk & Controls",
    "KPI, Portfolio & Decision Support",
]

ROLES = [
    {
        "title": "IT Consultant",
        "employer": "SafewayTax",
        "location": "Irvine, CA",
        "dates": "May 2025 - Aug 2026",
        "scope": "Hands-on technology modernization, cybersecurity, cloud operations, service delivery, and adoption leadership for a 75+ user professional services organization.",
        "bullets": [
            "Improved operational efficiency by 30% and workforce productivity by 20% by standardizing technology workflows, support routines, and user enablement.",
            "Migrated 85 people from Google Workspace to Microsoft 365 by directing planning, identity and data transition, testing, communications, training, and adoption.",
            "Sustained 99.9% infrastructure uptime and reduced downtime by 25% by strengthening operational controls, service ownership, vendor coordination, and escalation.",
            "Reduced security risk by 40% by improving identity, endpoint, access, monitoring, and prioritized remediation practices.",
            "Accelerated onboarding by 50% and reduced repeat support tickets by 45% by standardizing provisioning, documentation, self-service guidance, and root-cause resolution.",
            "Aligned modernization priorities to business needs by balancing security, reliability, usability, adoption, and small-organization capacity.",
        ],
    },
    {
        "title": "VP of IT, Promoted from Senior Director of IT",
        "employer": "Repipe Specialists",
        "location": "Burbank, CA",
        "dates": "Jun 2022 - Sep 2024",
        "scope": "Sole internal IT leader with enterprise accountability who built the operating model, then selected and governed an MSP as company headcount grew from 50+ to 100+.",
        "bullets": [
            "Supported 30% year-over-year business growth by translating business priorities into a scalable technology operating model, governance routines, and expanded service capacity.",
            "Reduced system downtime by 30% by strengthening infrastructure operations, service ownership, vendor accountability, and platform management practices.",
            "Lowered vulnerability exposure by 40% by formalizing cybersecurity controls, remediation discipline, identity and endpoint practices, and risk ownership.",
            "Cut incident resolution time by 35% by designing an MSP operating model with service expectations, escalation paths, performance reviews, and corrective actions.",
            "Improved executive decision speed by 25% by strengthening operational reporting, technology visibility, and leadership decision support.",
            "Improved workflow efficiency by 20% by leading NetSuite as business owner, directing third-party integrations, and administering the environment.",
            "Established portfolio and vendor governance by clarifying priorities, ownership, risks, dependencies, delivery expectations, and executive visibility.",
        ],
    },
    {
        "title": "IT End User Computing Leadership Assignment",
        "employer": "AVEVA",
        "location": "Lake Forest, CA",
        "dates": "Feb 2022 - May 2022",
        "scope": "Global end-user computing coordination in a distributed enterprise software environment.",
        "bullets": [
            "Supported approximately 2,800 employees and endpoints by coordinating a 10-person global end-user computing team without formal management authority.",
            "Maintained cross-region service alignment by clarifying ownership, coordinating priorities and handoffs, and reinforcing consistent endpoint and support practices.",
        ],
    },
    {
        "title": "Site IT Operations Leadership Assignment",
        "employer": "Cherokee Federal",
        "location": "Pomona, CA",
        "dates": "Aug 2021 - Dec 2021",
        "scope": "Mission-critical 24/7 technology operations with direct people leadership across three shifts.",
        "bullets": [
            "Sustained around-the-clock service coverage by leading seven direct reports across three shifts and strengthening ownership, handoffs, and escalation.",
            "Supported 2,400 employees and computers, approximately 1,000 mobile devices, and 650 additional tablets through disciplined endpoint and service operations.",
        ],
    },
    {
        "title": "Independent IT Consultant / Project Manager",
        "employer": "PTS",
        "location": "Huntington Beach, CA",
        "dates": "Aug 2019 - Jul 2021",
        "scope": "One-person consulting and project-delivery practice serving small and midsize businesses.",
        "bullets": [
            "Delivered practical infrastructure, business systems, security, and modernization initiatives by aligning each engagement to client priorities, risk, and operating capacity.",
            "Maintained end-to-end accountability by coordinating discovery, requirements, vendors, implementation, change, and stakeholder communication.",
        ],
    },
    {
        "title": "IT Director / Program Manager",
        "employer": "Master Engineering Services",
        "location": "Doha, Qatar",
        "dates": "Nov 2017 - Jun 2019",
        "scope": "Technology leadership for an engineering organization of approximately 250 employees.",
        "bullets": [
            "Improved delivery ownership by leading seven internal staff plus vendors across service delivery, infrastructure, business systems, and partner coordination.",
            "Strengthened operating consistency by aligning internal resources, external partners, technology priorities, and service expectations with business needs.",
        ],
    },
    {
        "title": "IT Operations Manager / Group IT Leader",
        "employer": "Jaidah Group",
        "location": "Doha, Qatar",
        "dates": "Mar 2006 - Oct 2017",
        "scope": "Enterprise technology leadership for 1,250 employees across 17 subsidiaries through direct and indirect reporting lines.",
        "bullets": [
            "Expanded technology capacity by growing the IT team from 9 to 17 while supporting 1,250 employees across a diversified 17-subsidiary group.",
            "Led an approximately $2M SAP S/4HANA program across 15 non-automotive subsidiaries and 650 employees by directing business alignment, partner delivery, process standardization, and adoption.",
            "Standardized automotive operations and reporting by delivering separate Incadea implementations for two automotive subsidiaries serving approximately 600 employees.",
            "Achieved 99.99% service availability by building and migrating a resilient multi-tenant data center and establishing operational ownership before cutover.",
            "Governed infrastructure, enterprise applications, service delivery, vendors, continuity, and business alignment across a complex multi-entity environment.",
        ],
    },
]

PROJECTS = [
    ("Scalable IT Operating Model", "Supported 30% year-over-year growth at Repipe Specialists by establishing governance, service ownership, cybersecurity discipline, vendor accountability, and a governed MSP model."),
    ("Business-Owned NetSuite Transformation", "Improved workflow efficiency by 20% by leading the implementation as business owner, personally directing third-party integrations, and administering the production environment."),
    ("Microsoft 365 Modernization", "Migrated 85 people from Google Workspace to Microsoft 365 while improving productivity by 20%, accelerating onboarding by 50%, and reducing repeat support tickets by 45%."),
    ("SAP S/4HANA and Incadea Portfolio", "Led an approximately $2M SAP program across 15 non-automotive subsidiaries and 650 employees, plus separate Incadea implementations for two automotive subsidiaries."),
    ("Multi-Tenant Data Center Build and Migration", "Achieved 99.99% service availability by governing design, migration, validation, operational readiness, and partner coordination across a 17-subsidiary group."),
    ("End-User Operations at Scale", "Coordinated service for approximately 2,800 employees and endpoints at AVEVA, and led seven direct reports supporting a large 24/7 device estate at Cherokee Federal."),
]


def rgb(value):
    return RGBColor.from_string(value)


def keep(paragraph, next_paragraph=False):
    props = paragraph._p.get_or_add_pPr()
    props.append(OxmlElement("w:keepLines"))
    if next_paragraph:
        props.append(OxmlElement("w:keepNext"))


def bottom_border(paragraph):
    props = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    border = OxmlElement("w:bottom")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), "8")
    border.set(qn("w:space"), "2")
    border.set(qn("w:color"), BLUE)
    borders.append(border)
    props.append(borders)


def set_cell_margins(cell, top=0, start=25, bottom=0, end=25):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def configure_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(9.4)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(2.3)
    normal.paragraph_format.line_spacing = 1.02

    title = document.styles["Title"]
    title.font.name = "Liberation Sans"
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = rgb(BLUE)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(1)

    heading = document.styles["Heading 1"]
    heading.font.name = "Liberation Sans"
    heading.font.size = Pt(13)
    heading.font.bold = True
    heading.font.color.rgb = rgb(BLUE)
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(3)
    heading.paragraph_format.keep_with_next = True


def section_heading(document, text):
    p = document.add_heading(text.upper(), level=1)
    bottom_border(p)
    keep(p, True)
    return p


def add_bullet(document, text, size=9.2):
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.30)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(1.3)
    run = p.add_run(text)
    run.font.size = Pt(size)
    keep(p)


def add_role(document, role, compact=False):
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(5.85)
    table.columns[1].width = Inches(1.35)
    for cell in table.rows[0].cells:
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP

    left = table.cell(0, 0).paragraphs[0]
    left.paragraph_format.space_before = Pt(4 if not compact else 2)
    left.paragraph_format.space_after = Pt(0)
    title = left.add_run(role["title"])
    title.bold = True
    title.font.size = Pt(10.2 if not compact else 9.5)
    left.add_run(" | ")
    employer = left.add_run(f"{role['employer']}, {role['location']}")
    employer.font.color.rgb = rgb("324A67")
    employer.font.size = Pt(10.2 if not compact else 9.5)

    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.space_before = Pt(4 if not compact else 2)
    right.paragraph_format.space_after = Pt(0)
    date = right.add_run(role["dates"])
    date.italic = True
    date.font.size = Pt(8.7 if not compact else 8.2)
    date.font.color.rgb = rgb(MUTED)

    scope = document.add_paragraph()
    scope.paragraph_format.space_after = Pt(1.5)
    scope_run = scope.add_run(role["scope"])
    scope_run.italic = True
    scope_run.font.size = Pt(8.7 if not compact else 8.2)
    scope_run.font.color.rgb = rgb(MUTED)
    keep(scope, True)
    for bullet in role["bullets"]:
        add_bullet(document, bullet, 8.9 if compact else 9.2)


def add_labeled_line(document, label, text):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(1.4)
    lead = p.add_run(f"{label}: ")
    lead.bold = True
    p.add_run(text)
    keep(p)


def build_docx(path):
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.66)
    section.right_margin = Inches(0.66)
    configure_styles(document)

    document.core_properties.title = "Joseph Hanna Executive Resume"
    document.core_properties.author = "Joseph Hanna"
    document.core_properties.subject = "Senior IT Executive, VP of IT, Head of IT, Director of IT, and CIO-track resume"
    document.core_properties.keywords = "VP of IT, CIO, Head of IT, Director of IT, technology strategy, cybersecurity, transformation"

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Joseph Hanna | Senior IT Executive | ")
    footer_run.font.name = "Liberation Sans"
    footer_run.font.size = Pt(7)
    footer_run.font.color.rgb = rgb(MUTED)
    page_number(footer)

    document.add_heading("JOSEPH HANNA", level=0)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    title_run = title.add_run("Senior IT Executive | VP of IT | Head of IT | CIO-Track")
    title_run.bold = True
    title_run.font.size = Pt(14)
    title_run.font.color.rgb = rgb("2F3E55")
    contact = document.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(12)
    contact_run = contact.add_run("Orange County, CA | Open to remote and Southern California hybrid roles | linkedin.com/in/joseph-g-hanna")
    contact_run.font.size = Pt(9.4)
    bottom_border(contact)

    section_heading(document, "Professional Summary")
    document.add_paragraph(
        "Senior technology executive with more than 20 years of experience translating business priorities into secure, scalable, and practical technology operations across growing, global, and multi-entity organizations. Leadership spans enterprise IT strategy, digital transformation, cybersecurity and resilience, cloud and infrastructure, ERP and business applications, service delivery, vendor governance, team development, and organizational change."
    )
    document.add_paragraph(
        "Demonstrated business impact includes supporting 30% year-over-year growth, reducing downtime by 30%, lowering vulnerability exposure by 40%, accelerating incident resolution by 35%, improving executive decision speed by 25%, migrating 85 people to Microsoft 365, and delivering 99.99% service availability. Known for combining executive partnership with hands-on accountability and adoption-focused execution."
    )

    section_heading(document, "Core Skills")
    skills = document.add_table(rows=6, cols=2)
    skills.alignment = WD_TABLE_ALIGNMENT.LEFT
    skills.autofit = False
    for row_index, row in enumerate(skills.rows):
        for col_index, cell in enumerate(row.cells):
            cell.width = Inches(3.55)
            set_cell_margins(cell, 0, 12, 0, 12)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0.8)
            p.add_run(f"• {CORE_SKILLS[col_index * 6 + row_index]}")

    section_heading(document, "Technical Skills")
    add_labeled_line(document, "Cloud & Infrastructure", "Microsoft 365, Azure, AWS, GCP, VMware, Hyper-V, modern workplace, shared infrastructure, endpoints, availability, DR/BCP")
    add_labeled_line(document, "Security & Governance", "Zero Trust, identity and access, Entra ID, Intune, Defender, Sentinel, Purview, SIEM, MFA, vulnerability remediation, risk and controls")
    add_labeled_line(document, "Enterprise Systems", "SAP S/4HANA, NetSuite, Incadea, Microsoft Dynamics CRM, Salesforce-connected environments, SharePoint, ERP/CRM integration")
    add_labeled_line(document, "Data, Automation & Reporting", "Power BI, executive KPI dashboards, portfolio reporting, workflow automation, PowerShell, SQL, data governance")
    add_labeled_line(document, "Operations & Delivery", "ITSM, incident leadership, MSP/MSSP governance, SLAs, vendor management, change management, adoption, service improvement")

    section_heading(document, "Work Experience")
    add_role(document, ROLES[0])

    document.add_page_break()
    section_heading(document, "Work Experience, Continued")
    for role in ROLES[1:]:
        add_role(document, role, compact=True)

    document.add_page_break()
    section_heading(document, "Education")
    education = [
        ("MBA, Project Management", "International Business School of Scandinavia"),
        ("MBA, Business Management", "Rome Business School"),
        ("Bachelor of Applied Science, Information Technology", "American Business and Technology University"),
        ("Bachelor of Commerce, Accounting", "Ain Shams University"),
    ]
    for degree, school in education:
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run(degree).bold = True
        p.add_run(f" - {school}")

    section_heading(document, "Certifications")
    certifications = [
        "Project Management Professional (PMP)",
        "ITIL Foundation",
        "Scrum Fundamentals Certified",
        "Microsoft Certified Technology Specialist (MCTS)",
        "Microsoft Certified Professional (MCP)",
        "Data Center Design Associate (DCDA)",
        "CIW 410 and CIW 430",
        "AI Fluency Framework & Foundations",
        "AI Fluency for Small Businesses",
        "Claude 101",
    ]
    for certification in certifications:
        add_bullet(document, certification, 9.4)

    section_heading(document, "Key Projects & Achievements")
    for name, detail in PROJECTS:
        p = document.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.30)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(f"{name}: ").bold = True
        p.add_run(detail)
        keep(p)

    document.save(path)


def main():
    soffice = shutil.which("soffice")
    if not soffice:
        raise SystemExit("LibreOffice is required to export the tagged resume PDF.")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    PUBLIC_OUT.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="resume-build-") as temp_dir:
        temp = Path(temp_dir)
        docx_path = temp / "Joseph_Hanna_Executive_Resume.docx"
        build_docx(docx_path)
        options = '{"UseTaggedPDF":{"type":"boolean","value":"true"},"PDFUACompliance":{"type":"boolean","value":"true"},"ExportBookmarks":{"type":"boolean","value":"true"}}'
        subprocess.run([
            soffice, "--headless", "--convert-to", f"pdf:writer_pdf_Export:{options}",
            "--outdir", str(temp), str(docx_path),
        ], check=True)
        generated = temp / "Joseph_Hanna_Executive_Resume.pdf"
        if not generated.exists():
            raise SystemExit("LibreOffice did not create the expected resume PDF.")
        shutil.copy2(generated, OUT)
        shutil.copy2(generated, PUBLIC_OUT)
    print(PUBLIC_OUT.relative_to(ROOT))


if __name__ == "__main__":
    main()
